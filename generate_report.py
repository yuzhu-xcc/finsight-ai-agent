"""
Week 2 / Day 12-14：Excel自动化报告生成
目标：跑一次脚本，自动产出一份带表格+走势图的Excel周报，
不用手动在Excel里一行行填数字、画图——这是"自动化"的意义所在。
"""
import sqlite3

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

# --- 第一步：跟之前一样，从数据库读数据、算指标 ---
conn = sqlite3.connect("finance.db")
df = pd.read_sql_query("SELECT * FROM stock_prices", conn)
conn.close()

df["date"] = pd.to_datetime(df["date"])
df = df.sort_values(["ticker", "date"])
df["daily_return_pct"] = df.groupby("ticker")["close"].pct_change() * 100

summary = df.groupby(["ticker", "name", "market"]).agg(
    start_price=("close", "first"),
    end_price=("close", "last"),
    volatility=("daily_return_pct", "std"),
).reset_index()
summary["total_return_pct"] = (
    (summary["end_price"] - summary["start_price"]) / summary["start_price"] * 100
)
summary = summary.round(2)

# --- 第二步：新建一个Excel工作簿 ---
# Workbook() 就是"新建一个空的Excel文件"，wb是变量名，代表这个文件
wb = Workbook()

# --- 第三步：制作"汇总"表 ---
# wb.active 拿到默认自带的第一张工作表，给它改个名字
ws_summary = wb.active
ws_summary.title = "股票汇总"

# 表头样式：加粗白字、深蓝底色，让表头一眼看出来是标题行
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

headers = ["股票代码", "名称", "市场", "期初价", "期末价", "总收益率(%)", "波动率"]
ws_summary.append(headers)
for cell in ws_summary[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# 把summary这张pandas表，一行行写进Excel
for _, row in summary.iterrows():
    ws_summary.append([
        row["ticker"], row["name"], row["market"],
        row["start_price"], row["end_price"],
        row["total_return_pct"], row["volatility"],
    ])

# 条件格式：涨的标绿、跌的标红——F列是"总收益率"这一列
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
for row_idx in range(2, ws_summary.max_row + 1):
    cell = ws_summary.cell(row=row_idx, column=6)  # 第6列 = 总收益率
    cell.fill = green_fill if cell.value >= 0 else red_fill

# 自动调整列宽，不然默认列宽太窄，数字会被截断显示
for col in ws_summary.columns:
    max_len = max(len(str(c.value)) for c in col) + 2
    ws_summary.column_dimensions[col[0].column_letter].width = max_len

# --- 第四步：给每支股票单独做一张"明细+走势图"的sheet ---
for ticker in summary["ticker"]:
    stock_df = df[df["ticker"] == ticker][["date", "close"]].copy()
    stock_df["date"] = stock_df["date"].dt.strftime("%Y-%m-%d")

    name = summary.loc[summary["ticker"] == ticker, "name"].values[0]
    ws = wb.create_sheet(title=f"{ticker}")

    ws.append(["日期", "收盘价"])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for _, row in stock_df.iterrows():
        ws.append([row["date"], row["close"]])

    # 画走势图：LineChart是"折线图"，Reference是"告诉图表去哪几列/哪几行取数据"
    chart = LineChart()
    chart.title = f"{name}（{ticker}）收盘价走势"
    chart.y_axis.title = "收盘价"
    chart.x_axis.title = "日期"

    data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 20
    chart.height = 10

    ws.add_chart(chart, "D2")  # "D2" = 图表放在D2这个单元格位置开始

# --- 第五步：保存文件 ---
wb.save("weekly_report.xlsx")
print("✅ 已生成 weekly_report.xlsx，在项目文件夹里应该能看到这个新文件了")
