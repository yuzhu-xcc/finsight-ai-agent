"""
FinSight 核心模块（终端版 agent.py 和网页版 streamlit_app.py 共用）

包含三部分：
1. 工具函数：查询股票表现、列出本地缓存、生成Excel报告
2. 系统提示词：规定AI怎么汇报数据
3. Agent对话循环：把"提问 -> 调用工具 -> 再回答"的流程封装成一个函数

为什么单独拆成一个文件：
之前 agent.py 和 streamlit_app.py 各抄了一份相同的工具代码，每改一处都要改两遍，很容易漏。
现在两个界面都只负责"收发消息"，核心逻辑只有这一份。
"""
import json
import os
import sqlite3
import time
from datetime import datetime, timedelta

import akshare as ak
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

MODEL = "claude-sonnet-5"
LOOKBACK_DAYS = 90        # 统一的滚动时间窗口：从"今天"往前推90天
MAX_REPORT_STOCKS = 10    # 一份Excel报告最多放几只股票
MAX_TOOL_ROUNDS = 8       # 单次提问最多连续调用几轮工具，防止死循环
DB_PATH = "finance.db"    # 本地缓存数据库（由 fetch_data.py 生成），只在实时查询失败时兜底


# ============================================================
# 第一部分：取数据（实时优先，本地缓存兜底）
# ============================================================

def _normalize_ticker(ticker) -> str:
    """统一股票代码格式：去掉空格、转大写、去掉 .HK/.SH/.SZ/.SS 这类后缀。"""
    t = str(ticker).strip().upper()
    for suffix in (".HK", ".SH", ".SZ", ".SS"):
        if t.endswith(suffix):
            t = t[: -len(suffix)]
    return t


def _fetch_live_series(ticker: str):
    """实时联网查询最近90天的日线收盘价。返回 (DataFrame[date, close], 市场)；失败则抛异常。"""
    end_date = datetime.now()
    start_date = end_date - timedelta(days=LOOKBACK_DAYS)

    if ticker.isdigit() and len(ticker) == 6:
        # A股6位代码：6开头是上海(sh)，其余按深圳(sz)处理
        prefix = "sh" if ticker.startswith("6") else "sz"
        df = ak.stock_zh_a_daily(symbol=f"{prefix}{ticker}", adjust="")
        df = df.reset_index()
        market = "A股"
    else:
        # 港股：补齐成5位数字（比如 700 -> 00700）
        df = ak.stock_hk_daily(symbol=ticker.zfill(5), adjust="")
        market = "港股"

    df["date"] = pd.to_datetime(df["date"]).dt.strftime("%Y-%m-%d")  # 统一成 YYYY-MM-DD，兼容各种日期格式
    df = df[(df["date"] >= start_date.strftime("%Y-%m-%d")) & (df["date"] <= end_date.strftime("%Y-%m-%d"))]
    df = df[["date", "close"]].sort_values("date").reset_index(drop=True)

    if df.empty:
        raise ValueError(f"没有找到股票代码 {ticker} 最近{LOOKBACK_DAYS}天的数据")
    return df, market


def _read_local(ticker: str):
    """从本地缓存数据库读某只股票的数据。没有数据库、没有这只股票、读取出错，一律返回None。"""
    if not os.path.exists(DB_PATH):
        return None
    try:
        conn = sqlite3.connect(DB_PATH)
        try:
            df = pd.read_sql_query(
                "SELECT date, close, name, market FROM stock_prices WHERE ticker = ?",
                conn,
                params=(ticker,),
            )
        finally:
            conn.close()
    except Exception:
        return None
    if df.empty:
        return None
    return df.sort_values("date").reset_index(drop=True)


def get_price_series(ticker):
    """
    取一只股票的价格序列：永远先实时联网查询；只有实时查询失败，才退回本地缓存。
    返回 (DataFrame[date, close], info字典)；实时和缓存都拿不到则抛 ValueError。
    """
    t = _normalize_ticker(ticker)

    try:
        df, market = _fetch_live_series(t)
    except Exception as live_error:
        cached = _read_local(t)
        if cached is None:
            raise ValueError(
                f"实时查询失败且本地也没有缓存数据：{live_error}（提示：目前只支持A股和港股，不支持美股）"
            ) from live_error
        return cached[["date", "close"]], {
            "ticker": t,
            "market": cached["market"].iloc[0],
            "name": cached["name"].iloc[0],
            "source": (
                f"⚠️本地缓存旧数据（{cached['date'].iloc[0]}至{cached['date'].iloc[-1]}，"
                "实时查询失败时的备用数据，可能不是最新）"
            ),
            "is_cached": True,
        }

    cached = _read_local(t)  # 实时数据本身不带公司名；如果本地缓存里有，就借用它的名字
    return df, {
        "ticker": t,
        "market": market,
        "name": cached["name"].iloc[0] if cached is not None else None,
        "source": f"实时查询（最近{LOOKBACK_DAYS}天）",
        "is_cached": False,
    }


def _summarize(df: pd.DataFrame, info: dict) -> dict:
    """把价格序列汇总成：起止日期、期初/期末价、总收益率、波动率。"""
    daily_return_pct = df["close"].pct_change() * 100
    start_price = df["close"].iloc[0]
    end_price = df["close"].iloc[-1]
    volatility = daily_return_pct.std()

    result = {
        "ticker": info["ticker"],
        "market": info["market"],
        "source": info["source"],
        "start_date": df["date"].iloc[0],
        "end_date": df["date"].iloc[-1],
        "start_price": round(float(start_price), 2),
        "end_price": round(float(end_price), 2),
        "total_return_pct": round(float((end_price - start_price) / start_price * 100), 2),
        "volatility": None if pd.isna(volatility) else round(float(volatility), 2),
    }
    if info.get("name"):
        result["name"] = info["name"]
    return result


# ============================================================
# 第二部分：工具函数（AI可以调用的"秘书技能"）
# ============================================================

def get_stock_summary(ticker: str) -> dict:
    """查询任意一支A股/港股最近90天的表现（实时优先，失败才用缓存）。"""
    try:
        df, info = get_price_series(ticker)
    except Exception as e:
        return {"error": str(e)}
    return _summarize(df, info)


def list_watchlist() -> list:
    """列出本地缓存数据库里有哪些股票（仅供实时查询失败时兜底参考）。"""
    if not os.path.exists(DB_PATH):
        return []
    try:
        conn = sqlite3.connect(DB_PATH)
        try:
            df = pd.read_sql_query("SELECT DISTINCT ticker, name, market FROM stock_prices", conn)
        finally:
            conn.close()
        return df.to_dict(orient="records")
    except Exception:
        return []


def _display_width(text) -> int:
    """估算文字在Excel里的显示宽度：中文字符占2格，其余占1格。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in str(text))


def generate_excel_report(stocks: list) -> dict:
    """
    生成Excel报告：一张"汇总"表 + 每只股票各一张"收盘价明细+走势图"。
    stocks 是列表，每项是 {"ticker": "00700", "name": "腾讯"}（name可省略，仅用于展示）。
    某只股票查询失败不会让整份报告失败：成功的写进报告，失败的单独列在返回结果里。
    """
    if not stocks:
        return {"error": "没有提供股票列表"}
    if len(stocks) > MAX_REPORT_STOCKS:
        return {"error": f"一份报告最多放{MAX_REPORT_STOCKS}只股票，请缩小范围后重试"}

    results, failed, seen = [], [], set()
    for item in stocks:
        raw_ticker = item.get("ticker", "") if isinstance(item, dict) else item
        given_name = item.get("name") if isinstance(item, dict) else None
        ticker = _normalize_ticker(raw_ticker)
        if not ticker or ticker in seen:
            continue
        seen.add(ticker)

        try:
            df, info = get_price_series(ticker)
        except Exception as e:
            failed.append({"ticker": ticker, "error": str(e)})
            continue

        summary = _summarize(df, info)
        # 名字优先级：本地缓存里的名字（可靠） > 调用方给的名字 > "—"
        summary["name"] = info.get("name") or given_name or "—"
        results.append((df, summary))
        time.sleep(0.5)  # 每只股票之间稍微停一下，减少被数据源限流的概率

    if not results:
        return {"error": "所有股票都查询失败，未生成报告", "failed": failed}

    # ---------- 样式 ----------
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    wb = Workbook()

    # ---------- 汇总表 ----------
    ws = wb.active
    ws.title = "汇总"
    headers = ["股票代码", "名称", "市场", "起始日期", "结束日期",
               "期初价", "期末价", "总收益率(%)", "波动率(%)", "数据来源"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for _, s in results:
        ws.append([s["ticker"], s["name"], s["market"], s["start_date"], s["end_date"],
                   s["start_price"], s["end_price"], s["total_return_pct"], s["volatility"], s["source"]])

    for row_idx in range(2, ws.max_row + 1):  # 总收益率那一列：涨绿跌红
        cell = ws.cell(row=row_idx, column=8)
        cell.fill = green_fill if (cell.value or 0) >= 0 else red_fill

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(_display_width(c.value) for c in col) + 2

    note_row = ws.max_row + 2
    ws.cell(row=note_row, column=1,
            value="说明：价格为日线收盘价；波动率 = 区间内每日收益率的标准差；"
                  f"报告生成时间 {datetime.now():%Y-%m-%d %H:%M}")

    # ---------- 每只股票一张明细 + 走势图 ----------
    for df, s in results:
        sheet = wb.create_sheet(title=s["ticker"][:31])
        sheet.append(["日期", "收盘价"])
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for date, close in zip(df["date"], df["close"]):
            sheet.append([date, float(close)])
        sheet.column_dimensions["A"].width = 14

        chart = LineChart()
        label = s["ticker"] if s["name"] == "—" else f"{s['name']}（{s['ticker']}）"
        chart.title = f"{label} 收盘价走势"
        chart.y_axis.title = "收盘价"
        chart.x_axis.title = "日期"
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=sheet.max_row), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row))
        chart.width, chart.height = 20, 10
        sheet.add_chart(chart, "D2")

    # ---------- 保存 ----------
    out_dir = os.environ.get("FINSIGHT_REPORT_DIR", "reports")
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"FinSight_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb.save(path)

    return {
        "file_path": path,
        "included": [
            {k: s.get(k) for k in ("ticker", "name", "market", "start_date", "end_date",
                                   "total_return_pct", "volatility", "source")}
            for _, s in results
        ],
        "failed": failed,
    }


TOOL_FUNCTIONS = {
    "get_stock_summary": get_stock_summary,
    "list_watchlist": list_watchlist,
    "generate_excel_report": generate_excel_report,
}

TOOLS = [
    {
        "name": "get_stock_summary",
        "description": (
            "查询任意一支A股或港股最近90天的整体表现（起止日期、期初价、期末价、总收益率、波动率）。"
            "永远优先实时联网查询，只有联网失败时才使用本地缓存的旧数据兜底（结果中会明确标注）。"
            "不支持美股。"
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "ticker": {"type": "string", "description": "股票代码，A股6位数字（如600519），港股4-5位数字（如00700或700）"},
            },
            "required": ["ticker"],
        },
    },
    {
        "name": "list_watchlist",
        "description": "列出本地缓存数据库里有数据的股票（仅在实时查询失败时作为兜底参考，不代表能查询的全部范围）。",
        "input_schema": {"type": "object", "properties": {}},
    },
    {
        "name": "generate_excel_report",
        "description": (
            "生成一份Excel报告文件：包含汇总表（收益率、波动率，涨绿跌红）和每只股票的收盘价明细+走势图，"
            f"数据为最近90天（实时优先）。当用户想要Excel、表格文件或报告时使用。一次最多{MAX_REPORT_STOCKS}只股票。"
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "stocks": {
                    "type": "array",
                    "description": "要放进报告的股票列表",
                    "items": {
                        "type": "object",
                        "properties": {
                            "ticker": {"type": "string", "description": "股票代码，A股6位数字，港股4-5位数字"},
                            "name": {"type": "string", "description": "可选，公司名称，仅用于报告展示，不参与计算"},
                        },
                        "required": ["ticker"],
                    },
                },
            },
            "required": ["stocks"],
        },
    },
]

SYSTEM_PROMPT = """你是 FinSight，一个专注于 A 股和港股数据分析的助手。

规则：
1. 你的主要专长是 A 股和港股数据分析；用户问其他话题时，可以正常简短回答，不需要拒绝。
2. 回答股价或表现时，必须明确告知数据对应的起止日期（工具返回的 start_date 和 end_date），并说明期末价格是该日期的收盘价，不要说成盘中实时价。
3. 如果数据来源标注为本地缓存旧数据，必须提醒用户这可能不是最新数据。
4. 不支持美股；用户问美股时，说明目前只支持 A 股和港股。
5. 用户想要 Excel / 表格文件 / 报告时，调用 generate_excel_report 工具生成。生成后简要说明报告包含哪些股票和数据日期；如果有股票失败（failed），如实告知是哪些、为什么；不要编造文件路径或下载链接。
6. 只陈述数据，不提供投资建议。
7. 用户用什么语言提问，就用什么语言回答。"""


# ============================================================
# 第三部分：Agent对话循环
# ============================================================

def extract_text(content_blocks) -> str:
    """从Claude的回复里找出真正的文字答案（回复里可能夹着"思考过程""工具调用"等其他类型的块）。"""
    for block in content_blocks:
        if block.type == "text":
            return block.text
    return ""


def run_agent_turn(client, messages, on_tool_result=None) -> str:
    """
    处理用户的一次提问：把对话历史发给Claude；它要调用工具就执行、把结果发回去，
    直到给出最终文字回答为止。messages 会被原地追加（保存完整对话历史）。
    on_tool_result(工具名, 参数, 结果)：可选回调，界面用它来打印过程或收集生成的文件。
    """
    for _ in range(MAX_TOOL_ROUNDS):
        response = client.messages.create(
            model=MODEL,
            max_tokens=2048,
            system=SYSTEM_PROMPT,
            tools=TOOLS,
            messages=messages,
        )
        messages.append({"role": "assistant", "content": response.content})

        if response.stop_reason != "tool_use":
            return extract_text(response.content)

        tool_results = []
        for block in response.content:
            if block.type != "tool_use":
                continue
            func = TOOL_FUNCTIONS.get(block.name)
            if func is None:
                result = {"error": f"未知工具：{block.name}"}
            else:
                try:
                    result = func(**block.input)
                except Exception as e:  # 任何工具内部出错，都变成"错误信息"交给AI解释，而不是让程序崩溃
                    result = {"error": f"工具执行出错：{e}"}
            if on_tool_result:
                on_tool_result(block.name, block.input, result)
            tool_results.append({
                "type": "tool_result",
                "tool_use_id": block.id,
                "content": json.dumps(result, ensure_ascii=False),
            })
        messages.append({"role": "user", "content": tool_results})

    return "（这个问题需要的工具调用次数过多，已自动中止，请换个更具体的问法。）"
